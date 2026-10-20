#!/usr/bin/env python3
"""Check that a vulnerability DOCX is self-contained enough for submission.

Usage:
  check-docx.py --docx report.docx --require "附录 A" --require "DeviceIoControl" \
                --forbid "见附件" --forbid "见 evidence" --min-images 4 --min-chars 5000
"""
import argparse
import json
import sys

from docx import Document


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", required=True)
    ap.add_argument("--require", action="append", default=[])
    ap.add_argument("--forbid", action="append", default=[])
    ap.add_argument("--min-images", type=int, default=0)
    ap.add_argument("--min-chars", type=int, default=0)
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    doc = Document(args.docx)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(parts)

    missing = [needle for needle in args.require if needle not in text]
    forbidden = [needle for needle in args.forbid if needle in text]
    images = len(doc.inline_shapes)
    chars = len(text)
    ok = not missing and not forbidden and images >= args.min_images and chars >= args.min_chars

    result = {
        "ok": ok,
        "docx": args.docx,
        "missing": missing,
        "forbidden_hits": forbidden,
        "images": images,
        "chars": chars,
        "min_images": args.min_images,
        "min_chars": args.min_chars,
    }
    if args.json:
        print(json.dumps(result, ensure_ascii=False))
    else:
        print(("OK" if ok else "FAIL") + f": {args.docx}")
        print(f"  chars={chars} images={images}")
        if missing:
            print("  missing: " + ", ".join(missing))
        if forbidden:
            print("  forbidden but present: " + ", ".join(forbidden))
    sys.exit(0 if ok else 2)


if __name__ == "__main__":
    main()
