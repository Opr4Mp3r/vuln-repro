#!/usr/bin/env python3
"""Convert a Markdown vulnerability report into a self-contained DOCX.

Usage:
  make_docx.py --md report-zh.md --out report.docx [--images images.json]
               [--shots-dir DIR] [--title TITLE] [--normalize-images]

images.json maps a level-2 section title (numbering stripped) to a list of
[filename, caption] pairs, e.g.:
  {"触发前提": [["03-winver.jpg", "图 3 ..."]],
   "复现步骤": [["07-stall-pre.jpg", "图 4 ..."]]}
"""
import argparse
import json
import os
import re
import sys
import tempfile

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMP_IMAGES = []


def set_cn_font(run, name="Microsoft YaHei", size=10.5, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def clean(text):
    return text.replace("`", "").replace("**", "").replace("__", "")


def add_para(doc, text, size=10.5, bold=False, mono=False, align=None,
             indent=True, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(clean(text))
    if mono:
        set_cn_font(run, "Consolas", 9, bold=False)
    else:
        set_cn_font(run, "Microsoft YaHei", size, bold=bold)
    return p


def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    set_cn_font(run, "Microsoft YaHei", 16 if level == 1 else 13, bold=True)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return h


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        for i in range(cols):
            txt = row[i] if i < len(row) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(clean(txt))
            set_cn_font(run, "Microsoft YaHei", 9.5)
    doc.add_paragraph()


def add_code(doc, code):
    for line in code:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        set_cn_font(run, "Consolas", 8.5)
    doc.add_paragraph()


def normalize_image(path):
    try:
        from PIL import Image
        img = Image.open(path).convert("RGB")
        fd, tmp = tempfile.mkstemp(suffix=".jpg")
        os.close(fd)
        img.save(tmp, "JPEG", quality=88)
        TEMP_IMAGES.append(tmp)
        return tmp
    except Exception as exc:
        print(f"warn: image normalization failed for {path}: {exc}", file=sys.stderr)
        return path


def add_image(doc, shots_dir, filename, caption, normalize=False):
    path = filename if os.path.isabs(filename) else os.path.join(shots_dir, filename)
    if not os.path.exists(path):
        print(f"warn: image not found: {path}", file=sys.stderr)
        return
    if normalize:
        path = normalize_image(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(15.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_cn_font(run, "Microsoft YaHei", 9)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    doc.add_paragraph()


def load_images(path):
    if not path:
        return {}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        out = {}
        for item in data:
            out.setdefault(item["section"], []).append([item["file"], item.get("caption", "")])
        return out
    return data


def convert(md_path, docx_path, images_by_section, shots_dir, title, normalize):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if title:
        doc.core_properties.title = title

    i = 0
    current_section = ""
    in_code = False
    code_buf = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                add_code(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.startswith("|"):
            table_buf = [line]
            j = i + 1
            while j < len(lines) and lines[j].startswith("|"):
                table_buf.append(lines[j])
                j += 1
            rows = []
            for tl in table_buf:
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                if all(set(c) <= set("-: ") for c in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            i = j
            continue
        if line.startswith("# "):
            text = title or clean(line[2:].strip())
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_cn_font(run, "Microsoft YaHei", 18, bold=True)
            doc.add_paragraph()
            current_section = ""
            i += 1
            continue
        if line.startswith("## "):
            if current_section in images_by_section:
                for fn, cap in images_by_section.pop(current_section):
                    add_image(doc, shots_dir, fn, cap, normalize)
            heading = clean(line[3:].strip())
            current_section = re.sub(r"^\d+\.\s*", "", heading)
            add_heading(doc, heading, 1)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, clean(line[4:].strip()), 2)
            i += 1
            continue
        if line.strip() == "":
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            # keep the list marker: ordered lists must not lose their numbers
            marker = m.group(2) if m.group(2).endswith(".") else "•"
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(f"{marker} {clean(m.group(3))}")
            set_cn_font(run, "Microsoft YaHei", 10.5)
            i += 1
            continue
        add_para(doc, line.strip())
        i += 1

    for section, imgs in images_by_section.items():
        for fn, cap in imgs:
            add_image(doc, shots_dir, fn, cap, normalize)

    doc.save(docx_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--images", default="")
    ap.add_argument("--shots-dir", default=".")
    ap.add_argument("--title", default="")
    ap.add_argument("--normalize-images", action="store_true")
    args = ap.parse_args()

    images = load_images(args.images)
    convert(args.md, args.out, images, args.shots_dir, args.title, args.normalize_images)
    for tmp in TEMP_IMAGES:
        try:
            os.unlink(tmp)
        except OSError:
            pass
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
